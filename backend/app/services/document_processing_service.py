"""Document Processing Service

Stage 4: Provide a single, reusable abstraction for converting user documents
(PDF/DOCX/XLSX/CSV/TXT/MD/JSON) into one or more artifact payloads.

Why:
- Explorer import, future upload endpoints, and agent tools should share the same
  conversion logic.
- Keeps explorer_service focused on filesystem browsing concerns.

This module is intentionally lightweight for now; later stages can improve
conversion quality (PDF cleanup, headings, etc.) without touching route code.

Public API:
- convert_file_to_artifact_payloads(file_path, *, imported_via="explorer")

Payload format:
- filename: str
- content: bytes
- artifact_meta: dict (optional)
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Optional

from .explorer_service import (
    convert_docx_to_md,
    convert_xlsx_to_md,
    convert_xlsx_to_json,
    convert_csv_to_md,
    convert_csv_to_json,
    convert_pdf_to_md,
    read_text_file,
)

logger = logging.getLogger(__name__)


def generate_docx_bytes(text: str) -> bytes:
    """Generate a simple .docx file from plain text.

    v1 scope (B1.7):
    - plain paragraphs
    - preserves blank lines as paragraph breaks
    """

    try:
        from docx import Document
    except ImportError as e:
        raise RuntimeError("python-docx not installed") from e

    doc = Document()
    # split on blank lines to preserve paragraph boundaries
    blocks = [b for b in (text or "").split("\n\n")]
    if not blocks:
        blocks = [""]
    for block in blocks:
        lines = block.splitlines() or [""]
        # keep manual line breaks within a paragraph
        para = doc.add_paragraph("\n".join(lines).rstrip())
        _ = para

    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_xlsx_bytes(
    *,
    sheets: list[dict],
) -> bytes:
    """Generate a simple .xlsx from JSON-like sheet payloads.

    Expected schema:
    sheets = [
      {
        "name": "Sheet1",
        "headers": ["A", "B"],
        "rows": [[1,2],[3,4]]
      }
    ]

    v1 scope (B1.7):
    - writes headers + rows
    - no styling
    """

    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise RuntimeError("openpyxl not installed") from e

    if not sheets:
        sheets = [{"name": "Sheet1", "headers": [], "rows": []}]

    wb = Workbook()

    # openpyxl creates a default sheet; reuse it for the first payload
    first = True
    for sheet in sheets:
        name = (sheet.get("name") or "Sheet").strip()[:31] or "Sheet"
        headers = sheet.get("headers") or []
        rows = sheet.get("rows") or []

        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(title=name)

        if headers:
            ws.append([str(h) if h is not None else "" for h in headers])
        for row in rows:
            if not isinstance(row, list):
                raise ValueError("Each row must be a list")
            ws.append(["" if v is None else v for v in row])

    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def convert_file_to_artifact_payloads(
    file_path: str,
    *,
    imported_via: str = "explorer",
) -> list[dict]:
    """Convert a file into one or more artifact payloads.

    Notes:
    - Does NOT write to disk and does NOT touch the DB.
    - Caller is responsible for creating Artifact DB rows.
    """

    p = Path(file_path)
    if not p.exists() or not p.is_file():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = p.suffix.lower()
    base_name = p.stem

    meta_base = {
        "source_file": p.name,
        "source_path": str(p),
        "imported_via": imported_via,
        "converted_from": ext,
    }

    payloads: list[dict] = []

    if ext == ".docx":
        content, error = convert_docx_to_md(file_path)
        if error:
            raise RuntimeError(error)
        payloads.append(
            {
                "filename": f"{base_name}.md",
                "content": content.encode("utf-8"),
                "artifact_meta": {**meta_base, "conversion": "docx_to_md"},
            }
        )

    elif ext in [".xlsx", ".xls"]:
        md_content, md_error = convert_xlsx_to_md(file_path)
        json_data, json_error = convert_xlsx_to_json(file_path)

        if md_error and json_error:
            raise RuntimeError(md_error or json_error)

        if not md_error:
            payloads.append(
                {
                    "filename": f"{base_name}.md",
                    "content": md_content.encode("utf-8"),
                    "artifact_meta": {**meta_base, "conversion": "xlsx_to_md"},
                }
            )
        if not json_error:
            payloads.append(
                {
                    "filename": f"{base_name}.json",
                    "content": json.dumps(json_data, indent=2, ensure_ascii=False).encode(
                        "utf-8"
                    ),
                    "artifact_meta": {**meta_base, "conversion": "xlsx_to_json"},
                }
            )

    elif ext == ".csv":
        md_content, md_error = convert_csv_to_md(file_path)
        json_data, json_error = convert_csv_to_json(file_path)

        if md_error and json_error:
            raise RuntimeError(md_error or json_error)

        if not md_error:
            payloads.append(
                {
                    "filename": f"{base_name}.md",
                    "content": md_content.encode("utf-8"),
                    "artifact_meta": {**meta_base, "conversion": "csv_to_md"},
                }
            )
        if not json_error:
            payloads.append(
                {
                    "filename": f"{base_name}.json",
                    "content": json.dumps(json_data, indent=2, ensure_ascii=False).encode(
                        "utf-8"
                    ),
                    "artifact_meta": {**meta_base, "conversion": "csv_to_json"},
                }
            )

    elif ext == ".pdf":
        content, error = convert_pdf_to_md(file_path)
        if error:
            raise RuntimeError(error)
        payloads.append(
            {
                "filename": f"{base_name}.md",
                "content": content.encode("utf-8"),
                "artifact_meta": {**meta_base, "conversion": "pdf_to_md"},
            }
        )

    elif ext in [".md", ".txt"]:
        content, error = read_text_file(file_path)
        if error:
            raise RuntimeError(error)
        target_name = f"{base_name}.md" if ext != ".md" else p.name
        payloads.append(
            {
                "filename": target_name,
                "content": content.encode("utf-8"),
                "artifact_meta": {**meta_base, "conversion": "text_to_md"},
            }
        )

    elif ext == ".json":
        content, error = read_text_file(file_path)
        if error:
            raise RuntimeError(error)
        payloads.append(
            {
                "filename": p.name,
                "content": content.encode("utf-8"),
                "artifact_meta": {**meta_base, "conversion": "copy_json"},
            }
        )

    else:
        raise RuntimeError(f"Unsupported file type: {ext}")

    return payloads
